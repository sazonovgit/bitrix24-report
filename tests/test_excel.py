from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from report_generator.excel import write_report
from report_generator.models import ReportRow


def _row(**kwargs) -> ReportRow:
    data = dict(
        title='Модуль "Витрина данных ГИСОГД"',
        year=2026,
        deal_id=479,
        deal_name="Сделка факт 2",
        company="Курганская область",
        price_po=Decimal("3000000"),
        price_impl=Decimal("450000"),
        amount=Decimal("3450000"),
        plan_fact="Факт",
    )
    data.update(kwargs)
    return ReportRow(**data)


def test_write_report_headers_and_numbers(tmp_path: Path):
    path = write_report(
        tmp_path / "report.xlsx",
        [
            _row(),
            _row(
                company="Омская область",
                deal_id=None,
                plan_fact="План",
                price_impl=None,
            ),
        ],
        year=2026,
    )
    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    assert headers == [
        "Название",
        "Год",
        "Сделка",
        "Название сделки",
        "Компания",
        "ЦЕНА ПО 2026",
        "ЦЕНА ВНЕДРЕНИЯ (2026)",
        "СУММА (2026)",
        "План/Факт",
    ]
    assert sheet["C2"].value == 479
    assert sheet["C3"].value is None
    assert sheet["F2"].value == 3000000
    assert sheet["I2"].value == "Факт"
    assert sheet["I3"].value == "План"


def test_write_report_selected_columns(tmp_path: Path):
    path = write_report(
        tmp_path / "subset.xlsx",
        [_row()],
        year=2025,
        columns=["title", "year", "amount", "plan_fact"],
    )
    sheet = load_workbook(path).active
    assert [cell.value for cell in sheet[1]] == [
        "Название",
        "Год",
        "СУММА (2026)",
        "План/Факт",
    ]
    assert sheet["C2"].value == 3450000


def test_price_headers_always_2026_when_report_year_differs(tmp_path: Path):
    path = write_report(
        tmp_path / "headers.xlsx",
        [_row(year=2025)],
        year=2025,
    )
    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    assert headers[5] == "ЦЕНА ПО 2026"
    assert headers[6] == "ЦЕНА ВНЕДРЕНИЯ (2026)"
    assert headers[7] == "СУММА (2026)"


def test_deal_name_hyperlink(tmp_path: Path):
    path = write_report(
        tmp_path / "link.xlsx",
        [_row(deal_name="Курган — закупка")],
        year=2026,
        columns=["deal_name"],
        portal_domain="portal.bitrix24.ru",
    )
    sheet = load_workbook(path).active
    cell = sheet["A2"]
    assert cell.value == "Курган — закупка"
    assert cell.hyperlink.target == "https://portal.bitrix24.ru/crm/deal/details/479/"
