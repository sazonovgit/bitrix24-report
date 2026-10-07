"""Запись отчёта в .xlsx."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from report_generator.bitrix import deal_details_url
from report_generator.config import COLUMN_KEYS
from report_generator.models import MONEY_COLUMNS, ReportRow, column_header


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=11)
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
THIN = Border(
    left=Side(style="thin", color="BDD7EE"),
    right=Side(style="thin", color="BDD7EE"),
    top=Side(style="thin", color="BDD7EE"),
    bottom=Side(style="thin", color="BDD7EE"),
)
FACT_FILL = PatternFill("solid", fgColor="E2EFDA")
PLAN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _cell_value(row: ReportRow, key: str):
    value = row.value(key)
    if key == "deal_id":
        return value if value else None
    if key in MONEY_COLUMNS and isinstance(value, Decimal):
        if value == value.to_integral_value():
            return int(value)
        return float(value)
    return value


def _autosize(sheet: Worksheet, columns: list[str]) -> None:
    for index, key in enumerate(columns, start=1):
        letter = get_column_letter(index)
        max_len = 10
        for cell in sheet[letter]:
            text = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(text), 70))
        width = max_len + 3
        if key in MONEY_COLUMNS:
            width = max(width, 18)
        sheet.column_dimensions[letter].width = width


def _apply_deal_name_cell(cell, row: ReportRow, portal_domain: str) -> None:
    cell.value = row.deal_name or None
    url = ""
    if row.deal_id and portal_domain:
        url = deal_details_url(portal_domain, int(row.deal_id))
    if url:
        cell.hyperlink = url
        cell.font = LINK_FONT
    else:
        cell.font = BODY_FONT


def write_report(
    path: Path,
    rows: list[ReportRow],
    *,
    year: int,
    columns: list[str] | None = None,
    portal_domain: str = "",
) -> Path:
    selected = [key for key in (columns or list(COLUMN_KEYS)) if key in COLUMN_KEYS]
    if not selected:
        raise ValueError("Выберите хотя бы одну колонку отчёта.")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"

    for col, key in enumerate(selected, start=1):
        cell = sheet.cell(1, col, column_header(key, year))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN

    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(selected))}{max(len(rows) + 1, 1)}"

    for row_index, row in enumerate(rows, start=2):
        fill = FACT_FILL if row.plan_fact == "Факт" else PLAN_FILL
        for col, key in enumerate(selected, start=1):
            if key == "deal_name":
                cell = sheet.cell(row_index, col)
                _apply_deal_name_cell(cell, row, portal_domain)
            else:
                cell = sheet.cell(row_index, col, _cell_value(row, key))
                cell.font = BODY_FONT
            cell.border = THIN
            cell.fill = fill
            if key in MONEY_COLUMNS:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif key in {"year", "deal_id"}:
                cell.alignment = Alignment(horizontal="center")
            elif key == "deal_name":
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=key == "title")

    _autosize(sheet, selected)
    workbook.save(path)
    return path
